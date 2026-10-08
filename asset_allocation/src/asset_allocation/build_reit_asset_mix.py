"""Validate the supplied Codal asset-mix workbook and publish a derived chart panel."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path

import jdatetime
import pandas as pd
from openpyxl import load_workbook

from asset_allocation.analyze_reit_tedpix_weekly import atomic_frame
from asset_allocation.build_reit_reinvested_correlations import previous_month

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/interim/codal_reit_asset_mix/REIT_Asset_Mix_24M.xlsx"
OUTPUT = ROOT / "data/processed/analysis/reit_codal_asset_mix_monthly.csv"
SOURCE_SHA256 = "a2592e41fee12f28db32566224794d3a812969f1405845e852270d58f73506a1"
FUNDS = {"ارزش مسکن": "Arzesh Maskan", "دانیک": "Danik"}
COMPONENTS = (
    "Direct Real Estate (bn IRR)", "Fixed-Income Funds (bn IRR)",
    "Direct Fixed-Income Securities (bn IRR)", "Direct Equity (bn IRR)",
    "Cash & Bank (bn IRR)", "Other / Receivables (bn IRR)",
)


def extract(source: Path = SOURCE, *, expected_hash: str = SOURCE_SHA256) -> pd.DataFrame:
    digest = sha256(source.read_bytes()).hexdigest()
    if digest != expected_hash:
        raise ValueError("Codal asset-mix snapshot changed; preserve it and review the new source separately")
    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        sheet = workbook["Asset_Mix_24M"]
        records = sheet.iter_rows(values_only=True)
        headers = next(records)
        if len(headers) != len(set(headers)):
            raise ValueError("Duplicate asset-mix workbook column")
        rows = []
        for record in records:
            values = dict(zip(headers, record))
            if not values.get("Fund"):
                continue
            if values["Fund"] not in FUNDS:
                raise ValueError(f"Unexpected fund: {values['Fund']}")
            year, month, day = map(int, str(values["Jalali Date"]).split("/"))
            observed = jdatetime.date(year, month, day).togregorian()
            period = f"{year:04}/{month:02}"
            amounts = [float(values[column]) for column in COMPONENTS]
            if any(amount < 0 for amount in amounts):
                raise ValueError(f"Negative component: {values['Fund']} {period}")
            total = sum(amounts)
            if total <= 0 or abs(total - float(values["Total Assets (bn IRR)"])) > 1e-5:
                raise ValueError(f"Asset components do not reconcile: {values['Fund']} {period}")
            housing, fixed_funds, fixed_direct, equity, cash, receivables = amounts
            reported = (values["Real Estate %"], values["Total Fixed-Income %"],
                        values["Equity %"], values["Cash %"], values["Other %"])
            computed = (housing / total, (fixed_funds + fixed_direct) / total,
                        equity / total, cash / total, receivables / total)
            if any(value is None or abs(float(value) - actual) > 1e-6
                   for value, actual in zip(reported, computed)):
                raise ValueError(f"Workbook percentage mismatch: {values['Fund']} {period}")
            rows.append({"fund": FUNDS[values["Fund"]], "jalali_period": period,
                         "source_jalali_date": values["Jalali Date"],
                         "observation_date_gregorian": observed.isoformat(),
                         "housing_share": computed[0], "fixed_income_share": computed[1],
                         "other_including_cash_equity_share": sum(computed[2:]),
                         "equity_share": computed[2], "cash_share": computed[3],
                         "receivables_share": computed[4],
                         "total_assets_bn_irr": total,
                         "reconstruction_method": values["Reconstruction Method"],
                         "data_quality_flag": values["Data Quality Flag"] or "",
                         "source_sha256": digest})
    finally:
        workbook.close()
    panel = pd.DataFrame(rows).sort_values(["fund", "jalali_period"]).reset_index(drop=True)
    if set(panel.fund) != set(FUNDS.values()) or panel.duplicated(["fund", "jalali_period"]).any():
        raise ValueError("Incomplete or duplicate fund/month asset mix")
    for fund, sample in panel.groupby("fund"):
        if len(sample) != 24 or any(previous_month(current) != prior for prior, current in
                                    zip(sample.jalali_period.iloc[:-1], sample.jalali_period.iloc[1:])):
            raise ValueError(f"Expected 24 consecutive months for {fund}")
        if (sample[["housing_share", "fixed_income_share", "other_including_cash_equity_share"]]
                .sum(axis=1).sub(1).abs() > 1e-8).any():
            raise ValueError(f"Asset mix does not sum to 100% for {fund}")
    return panel


def build(source: Path = SOURCE) -> pd.DataFrame:
    panel = extract(source)
    atomic_frame(OUTPUT, panel)
    return panel


if __name__ == "__main__":
    result = build()
    print(f"Wrote {len(result)} fund/month allocation rows through {result.jalali_period.max()}")
