"""Export four source-traceable REIT data files for presentation authoring."""
from __future__ import annotations

import json
import tempfile
from datetime import date, timedelta
from pathlib import Path

import jdatetime
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from asset_allocation.analyze_reit_tedpix_weekly import atomic_frame, shift_jalali_months
from asset_allocation.build_reit_assembly_reinvestment import FUNDS, ROOT

RAW = ROOT / "data/raw/real_estate_funds"
DATA = ROOT / "data/processed/analysis"
POST_WAR_START = date(2026, 4, 9)  # Day after 40 inclusive days from 2026-02-28.
ANCHOR = date(2026, 10, 9)


def jalali(value: str) -> str:
    d = jdatetime.date.fromgregorian(date=date.fromisoformat(value))
    return f"{d.year:04}/{d.month:02}/{d.day:02}"


def trading_value_archive() -> dict[tuple[str, str], dict]:
    values = {}
    for path in sorted((RAW / "snapshots").glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        for row in payload.get("closingPriceDaily", []):
            code = str(row["insCode"])
            day = str(row["dEven"])
            key = code, f"{day[:4]}-{day[4:6]}-{day[6:]}"
            value = int(row["qTotCap"])
            close = float(row["pClosing"])
            if value < 0 or close <= 0:
                raise ValueError(f"Invalid TSETMC snapshot row {key}")
            old = values.get(key)
            if old and (old["trading_value_irr"] != value or old["snapshot_close"] != close):
                raise ValueError(f"Conflicting immutable snapshot values {key}")
            values[key] = {"trading_value_irr": value, "snapshot_close": close,
                           "trading_value_snapshot": path.stem}
    return values


def window_return(group: pd.DataFrame, cutoff: date, level: str, max_gap_days: int | None = 14) -> tuple[str | None, float | None]:
    sample = group.loc[group.source_date_gregorian.ge(cutoff.isoformat())]
    if len(sample) < 2:
        return None, None
    first = sample.iloc[0]
    if max_gap_days is not None and (date.fromisoformat(first.source_date_gregorian) - cutoff).days > max_gap_days:
        return None, None
    return first.source_date_gregorian, float(sample[level].iloc[-1] / first[level] - 1)


def build_daily() -> pd.DataFrame:
    archive = trading_value_archive()
    panel = pd.read_csv(DATA / "reit_assembly_reinvested_daily.csv",
                        dtype={"ins_code": str, "source_date_gregorian": str})
    if panel.duplicated(["ins_code", "source_date_gregorian"]).any():
        raise ValueError("Duplicate derived fund/date")
    manifest = pd.read_csv(RAW / "instruments.csv", dtype={"ins_code": str})
    if set(manifest.ins_code) != set(FUNDS):
        raise ValueError("Raw instrument roster does not match eight configured funds")
    rows = []
    for code, name in FUNDS.items():
        raw = pd.read_csv(RAW / f"{code}.csv", dtype={"ins_code": str, "source_date_gregorian": str})
        if raw.duplicated("source_date_gregorian").any() or not raw.ins_code.eq(code).all():
            raise ValueError(f"Invalid raw instrument mapping for {name}")
        fund_panel = panel.loc[panel.ins_code.eq(code)].set_index("source_date_gregorian")
        for r in raw.itertuples():
            key = (code, r.source_date_gregorian)
            if key not in archive:
                raise ValueError(f"Missing qTotCap snapshot for {key}")
            source = archive[key]
            if float(r.closing_price_irr) != source["snapshot_close"]:
                raise ValueError(f"Snapshot/raw closing price mismatch {key}")
            trade = bool(r.has_trade)
            if trade != (r.source_date_gregorian in fund_panel.index):
                raise ValueError(f"Traded-session mismatch {key}")
            derived = fund_panel.loc[r.source_date_gregorian] if trade else None
            rows.append({
                "fund": name, "fund_persian": r.ticker, "ins_code": code,
                "date_gregorian": r.source_date_gregorian,
                "date_jalali": jalali(r.source_date_gregorian),
                "raw_close_irr": r.closing_price_irr,
                "trade_count": r.trade_count, "trade_volume_units": r.trade_volume,
                "has_trade": trade,
                "trading_value_irr_qTotCap": source["trading_value_irr"],
                "trading_value_billion_toman": source["trading_value_irr"] / 1e10,
                "price_return_index": derived.price_index if trade else None,
                "hypothetical_total_return_index": derived.total_return_index if trade else None,
                "reinvested_units": derived.reinvested_units if trade else None,
                "dividend_receivable_irr": derived.dividend_receivable_irr if trade else None,
                "qTotCap_snapshot_sha256": source["trading_value_snapshot"],
            })
    daily = pd.DataFrame(rows).sort_values(["ins_code", "date_gregorian"]).reset_index(drop=True)
    atomic_frame(DATA / "reit_daily_series.csv", daily)
    return daily


def build_profiles(daily: pd.DataFrame) -> pd.DataFrame:
    panel = pd.read_csv(DATA / "reit_assembly_reinvested_daily.csv", dtype={"ins_code": str})
    events = pd.read_csv(DATA / "reit_dividend_events_verified.csv", dtype={"ins_code": str})
    official = pd.read_csv(DATA / "reit_tsetmc_instrument_info.csv", dtype={"ins_code": str}).set_index("ins_code")
    subscriptions = pd.read_csv(DATA / "reit_subscription_facts.csv", dtype={"ins_code": str}).set_index("ins_code")
    if set(official.index) != set(FUNDS) or official.index.has_duplicates:
        raise ValueError("Official REIT profile source table has a mismatched or duplicate roster")
    if set(subscriptions.index) != set(FUNDS) or subscriptions.index.has_duplicates:
        raise ValueError("REIT subscription source table has a mismatched or duplicate roster")
    result = []
    for code, name in FUNDS.items():
        raw = daily.loc[daily.ins_code.eq(code)]
        traded = raw.loc[raw.has_trade.eq(True)].sort_values("date_gregorian")
        fund_panel = panel.loc[panel.ins_code.eq(code)].sort_values("source_date_gregorian")
        first, last = traded.iloc[0], traded.iloc[-1]
        last_day = date.fromisoformat(last.date_gregorian)
        one_year_cutoff = last_day - timedelta(days=365)
        post = traded.loc[traded.date_gregorian.ge(POST_WAR_START.isoformat())]
        year = traded.loc[traded.date_gregorian.ge(one_year_cutoff.isoformat())]
        post_base, post_price = window_return(fund_panel, POST_WAR_START, "traded_close_irr", max_gap_days=None)
        _, post_total = window_return(fund_panel, POST_WAR_START, "reinvested_value_irr", max_gap_days=None)
        if date.fromisoformat(first.date_gregorian) > one_year_cutoff:
            year_base = year_price = year_total = None
        else:
            year_base, year_price = window_return(fund_panel, one_year_cutoff, "traded_close_irr")
            _, year_total = window_return(fund_panel, one_year_cutoff, "reinvested_value_irr")
        fund_events = events.loc[events.ins_code.eq(code)]
        sourced = official.loc[code]
        subscription = subscriptions.loc[code]
        if sourced.latest_close_date_gregorian != last.date_gregorian or float(sourced.latest_close_irr) != float(last.raw_close_irr):
            raise ValueError(f"Official profile source has stale close for {code}")
        same_date = bool(sourced.same_date_market_cap_verified)
        result.append({
            "fund": name, "fund_persian": last.fund_persian, "ins_code": code,
            "first_trading_date_gregorian": first.date_gregorian,
            "first_trading_date_jalali": first.date_jalali,
            "subscription_date_jalali": subscription.subscription_date_jalali,
            "subscription_proceeds_billion_toman": subscription.subscription_proceeds_billion_toman,
            "founder_or_sponsor": subscription.founder_or_sponsor,
            "founder_status": subscription.founder_status,
            "subscription_notice_url": subscription.subscription_notice_url,
            "founder_source_url": subscription.founder_source_url,
            "official_nav_billion_toman": sourced.official_nav_billion_toman,
            "official_nav_date_jalali": sourced.official_nav_date_jalali,
            "official_nav_source_url": sourced.fipiran_source_url,
            "latest_reported_units": sourced.tsetmc_units_zTitad,
            "latest_reported_units_date_gregorian": sourced.tsetmc_instrument_date_gregorian,
            "units_outstanding_asof_close": sourced.tsetmc_units_zTitad if same_date else None,
            "market_cap_billion_toman": sourced.indicative_cap_with_latest_close_billion_toman if same_date else None,
            "indicative_cap_with_latest_close_billion_toman": sourced.indicative_cap_with_latest_close_billion_toman,
            "latest_reported_units_source_url": sourced.source_url,
            "latest_close_irr": last.raw_close_irr,
            "latest_close_date_gregorian": last.date_gregorian,
            "latest_close_date_jalali": last.date_jalali,
            "one_year_turnover_start_gregorian": one_year_cutoff.isoformat(),
            "one_year_traded_sessions": len(year),
            "avg_daily_trading_value_one_year_billion_toman": year.trading_value_billion_toman.mean(),
            "post_war_start_gregorian": POST_WAR_START.isoformat(),
            "post_war_start_jalali": jalali(POST_WAR_START.isoformat()),
            "post_war_first_traded_date": post_base,
            "post_war_history_status": "listed_after_post_war_start" if date.fromisoformat(first.date_gregorian) > POST_WAR_START else "listed_before_post_war_start",
            "post_war_traded_sessions": len(post),
            "avg_daily_trading_value_post_war_billion_toman": post.trading_value_billion_toman.mean() if len(post) else None,
            "avg_daily_trading_value_since_first_trade_billion_toman": traded.trading_value_billion_toman.mean(),
            "post_war_price_return": post_price,
            "post_war_hypothetical_total_return": post_total,
            "one_year_baseline_date": year_base,
            "one_year_price_return": year_price,
            "one_year_hypothetical_total_return": year_total,
            "since_first_trade_price_return": fund_panel.traded_close_irr.iloc[-1] / fund_panel.traded_close_irr.iloc[0] - 1,
            "since_first_trade_hypothetical_total_return": fund_panel.reinvested_value_irr.iloc[-1] / fund_panel.reinvested_value_irr.iloc[0] - 1,
            "identified_positive_dividend_events": len(fund_events),
            "identified_dividends_per_original_unit_irr": fund_events.cash_irr_per_unit.sum(),
            "dividend_history_status": "provisional_dates_and_incomplete" if name == "Amin Shahr"
                                       else "unverified_dividend_history" if not len(fund_events)
                                       else "incomplete_dividend_history",
            "profile_missing_fields": "" if same_date else "units_outstanding_asof_close,market_cap_asof_close",
            "return_method": "raw_close_first_traded_session_after_event_reinvestment",
            "trading_value_method": "TSETMC_qTotCap_IRR_divided_by_1e10_traded_sessions_only",
        })
    profile = pd.DataFrame(result)
    atomic_frame(DATA / "reit_fund_profiles.csv", profile)
    return profile


def build_dividends() -> pd.DataFrame:
    events = pd.read_csv(DATA / "reit_dividend_events_verified.csv", dtype={"ins_code": str})
    columns = {
        "fund": "fund", "ins_code": "ins_code", "event_id": "event_id",
        "event_type": "event_type", "assembly_date_jalali": "assembly_date_jalali",
        "assembly_date_gregorian": "assembly_date_gregorian",
        "reported_distribution_date_jalali": "reported_distribution_date_jalali",
        "event_date_jalali": "model_event_date_jalali",
        "event_date_gregorian": "model_event_date_gregorian",
        "cash_irr_per_unit": "dividend_irr_per_eligible_unit",
        "eligibility_reference_event": "eligibility_reference_event",
        "reinvestment_date_gregorian": "hypothetical_reinvestment_date_gregorian",
        "reinvestment_price_irr": "raw_reinvestment_close_irr",
        "units_entitled": "units_entitled_from_one_initial_unit",
        "evidence_status": "date_evidence_status",
        "data_source": "source_description",
        "notes": "notes",
    }
    if "event_date_gregorian" not in events:
        events["event_date_gregorian"] = events["effective_model_event_date_gregorian"]
    manifest = pd.read_csv(RAW / "instruments.csv", dtype={"ins_code": str})
    persian = dict(zip(manifest.ins_code, manifest.ticker))
    events["fund_persian"] = events.ins_code.map(persian)
    events["hypothetical_reinvestment_date_jalali"] = events.reinvestment_date_gregorian.map(jalali)
    dividends = events[["fund", "fund_persian", "ins_code", "event_id", "event_type",
                        "assembly_date_jalali", "assembly_date_gregorian",
                        "reported_distribution_date_jalali", "event_date_jalali",
                        "event_date_gregorian", "cash_irr_per_unit",
                        "eligibility_reference_event", "reinvestment_date_gregorian",
                        "hypothetical_reinvestment_date_jalali", "reinvestment_price_irr",
                        "units_entitled", "evidence_status", "data_source", "notes"]].rename(columns=columns)
    atomic_frame(DATA / "reit_dividends.csv", dividends)
    return dividends


def add_sheet(book: Workbook, title: str, data: pd.DataFrame, percent_columns: set[str] | None = None):
    ws = book.create_sheet(title)
    ws.append(list(data.columns))
    percent_columns = percent_columns or set()
    for row in data.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="19334A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, field in enumerate(data.columns, 1):
        vals = data[field].astype(str)
        width = min(34, max(13, len(field) + 2, int(vals.str.len().quantile(.9)) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
        for col in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
            for cell in col:
                if (field in percent_columns or (title.endswith("_r") and field != "fund")) and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000"
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F5F8")
    return ws


def build_correlations() -> Path:
    benchmark = pd.read_csv(DATA / "reit_reinvested_correlations.csv")
    benchmark = benchmark.loc[benchmark.frequency.eq("weekly") &
                              benchmark.window_months.isin([3, 6, 12, 24]) &
                              benchmark.benchmark.isin(["TEDPIX", "USD/IRR"])].copy()
    benchmark["window_start_exclusive_gregorian"] = benchmark.window_months.map(
        lambda months: shift_jalali_months(ANCHOR, int(months)).isoformat())
    benchmark["window_end_inclusive_gregorian"] = ANCHOR.isoformat()
    bench_fields = ["fund", "window_months", "window_start_exclusive_gregorian",
                    "window_end_inclusive_gregorian", "paired_observations", "minimum_pairs",
                    "correlation", "status"]
    heat = pd.read_csv(DATA / "reit_requested_correlation_heatmaps.csv")
    housing = heat.loc[heat.benchmark.eq("Tehran housing") &
                       heat.frequency.eq("monthly") &
                       heat.window_months.isin([12, 24])].copy()
    housing["window_end_jalali_month"] = "1405/06"
    housing = housing[["fund", "window_months", "window_end_jalali_month",
                       "paired_observations", "minimum_pairs", "correlation", "status"]]
    weekly = pd.read_csv(DATA / "reit_all_fund_weekly_returns.csv")
    book = Workbook()
    book.remove(book.active)
    for source, title in [("TEDPIX", "TEDPIX"), ("USD/IRR", "USD_IRR")]:
        part = benchmark.loc[benchmark.benchmark.eq(source), bench_fields].sort_values(["fund", "window_months"])
        add_sheet(book, title, part.reset_index(drop=True), {"correlation"})
    add_sheet(book, "Tehran_housing", housing.sort_values(["fund", "window_months"]).reset_index(drop=True), {"correlation"})
    names = list(FUNDS.values())
    for months in (3, 12):
        cutoff = shift_jalali_months(ANCHOR, months).isoformat()
        sample = weekly.loc[weekly.period.gt(cutoff) & weekly.period.le(ANCHOR.isoformat())]
        wide = sample.pivot(index="period", columns="fund", values="total_return")
        coefficients, counts = [], []
        for left in names:
            c_row, n_row = {"fund": left}, {"fund": left}
            for right in names:
                pair = wide[[left]].dropna() if left == right else wide[[left, right]].dropna()
                n = len(pair)
                n_row[right] = n
                if n < 5 or pair[left].nunique() <= 1:
                    c_row[right] = None
                elif left == right:
                    c_row[right] = 1.0
                else:
                    c_row[right] = pair[left].corr(pair[right]) if pair[right].nunique() > 1 else None
            coefficients.append(c_row)
            counts.append(n_row)
        add_sheet(book, f"REIT_{months}m_r", pd.DataFrame(coefficients))
        add_sheet(book, f"REIT_{months}m_n", pd.DataFrame(counts))
    method = pd.DataFrame([
        ["market_date", "2026-10-07", "Latest valid REIT and TEDPIX traded close"],
        ["weekly_anchor", ANCHOR.isoformat(), "Completed Friday; last observed close may be Wednesday"],
        ["weekly_windows", "3, 6, 12, 24 Jalali months", "Start exclusive, end inclusive; Friday periods"],
        ["housing_end", "1405/06", "Monthly Tehran housing listing-price proxy"],
        ["return_method", "hypothetical assembly/event-date reinvestment", "Raw traded closes; first later valid session"],
        ["suppression", "benchmark project coverage rules; pairwise n >= 5", "Blank correlation is insufficient coverage or variation"],
        ["missing_prices", "no interpolation", "Only observed returns enter correlations"],
    ], columns=["item", "value", "meaning"])
    add_sheet(book, "Method", method)
    output = DATA / "reit_correlations.xlsx"
    with tempfile.NamedTemporaryFile(dir=DATA, suffix=".xlsx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def build() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, Path]:
    daily = build_daily()
    from asset_allocation.build_reit_subscription_facts import build as build_subscriptions
    build_subscriptions()
    profiles = build_profiles(daily)
    dividends = build_dividends()
    workbook = build_correlations()
    return profiles, daily, dividends, workbook


if __name__ == "__main__":
    profiles, daily, dividends, workbook = build()
    print(f"profiles={len(profiles)} daily={len(daily)} dividends={len(dividends)} workbook={workbook}")
